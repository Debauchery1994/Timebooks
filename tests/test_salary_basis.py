"""Tests for salaries entered per year or per pay period (Employee.salary_basis)."""
from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl

from timecard import calc
from timecard.models import Company, Employee, HoursSummary
from timecard.storage import TimecardBook

D = Decimal


def pay(emp: Employee, company: Company, hours: HoursSummary | None = None) -> Decimal:
    return calc.compute_payroll(emp, hours or HoursSummary(regular=D("40")), company, check_number=1,
                                check_date=date(2026, 1, 9), period_start=date(2026, 1, 5),
                                period_end=date(2026, 1, 11)).gross_pay


class SalaryBasisPayTests(unittest.TestCase):
    def test_per_period_salary_is_the_check_amount_on_every_schedule(self) -> None:
        emp = Employee("E001", "Ann", "Lee", pay_type="Salary", pay_rate=D("1250.00"), salary_basis="Per Period")
        for period in ("Weekly", "Biweekly", "Semimonthly", "Monthly"):
            with self.subTest(period=period):
                self.assertEqual(pay(emp, Company(pay_period=period)), D("1250.00"))

    def test_per_period_ignores_hours_and_overtime(self) -> None:
        emp = Employee("E001", "Ann", "Lee", pay_type="Salary", pay_rate=D("900"), salary_basis="Per Period")
        c = Company(overtime_enabled=True)
        self.assertEqual(pay(emp, c, HoursSummary(regular=D("40"), overtime=D("10"))), D("900.00"))
        self.assertEqual(pay(emp, c, HoursSummary()), D("900.00"))

    def test_annual_is_unchanged(self) -> None:
        emp = Employee("E001", "Ann", "Lee", pay_type="Salary", pay_rate=D("52000"))
        self.assertEqual(emp.salary_basis, "Annual")
        self.assertEqual(pay(emp, Company(pay_period="Biweekly")), D("2000.00"))

    def test_basis_is_ignored_for_hourly(self) -> None:
        emp = Employee("E001", "Ann", "Lee", pay_rate=D("20"), salary_basis="Per Period")
        self.assertEqual(pay(emp, Company()), D("800.00"))

    def test_basis_spellings(self) -> None:
        for text, per in (("Per Period", True), ("per period", True), ("PerPeriod", True),
                          ("Annual", False), ("", False)):
            emp = Employee("E1", "A", "B", pay_type="Salary", salary_basis=text)
            self.assertEqual(calc.salary_is_per_period(emp), per, text)


class SalaryBasisStorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.path = Path(self._tmp.name) / "book.xlsx"

    def test_round_trip(self) -> None:
        book = TimecardBook.new(Company(year=2026))
        book.upsert_employee(Employee("E001", "Ann", "Lee", pay_type="Salary", salary_basis="Per Period"))
        book.set_pay_rate("E001", D("1250"))
        book.upsert_employee(Employee("E002", "Bo", "Ray", pay_type="Salary"))
        book.upsert_employee(Employee("E003", "Cy", "Fox"))
        book.save(self.path)
        loaded = TimecardBook.load(self.path)
        self.assertEqual(loaded.employees["E001"].salary_basis, "Per Period")
        self.assertEqual(loaded.employees["E001"].pay_rate, D("1250.00"))
        self.assertEqual(loaded.employees["E002"].salary_basis, "Annual")
        ws = openpyxl.load_workbook(self.path)["Employees"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[-1], "SalaryBasis")
        col = headers.index("SalaryBasis")
        self.assertEqual([r[col] for r in ws.iter_rows(min_row=2, values_only=True)],
                         ["Per Period", "Annual", None])   # blank for hourly

    def test_upsert_keeps_rate_but_updates_basis(self) -> None:
        book = TimecardBook.new(Company(year=2026))
        book.upsert_employee(Employee("E001", "Ann", "Lee", pay_type="Salary"))
        book.set_pay_rate("E001", D("52000"))
        book.upsert_employee(Employee("E001", "Ann", "Lee", pay_type="Salary", salary_basis="Per Period"))
        self.assertEqual(book.employees["E001"].salary_basis, "Per Period")
        self.assertEqual(book.employees["E001"].pay_rate, D("52000.00"))

    def test_older_workbook_without_the_column_loads_as_annual(self) -> None:
        book = TimecardBook.new(Company(year=2026))
        book.upsert_employee(Employee("E001", "Ann", "Lee", pay_type="Salary"))
        book.set_pay_rate("E001", D("52000"))
        book.save(self.path)
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Employees"]
        ws.delete_cols([c.value for c in ws[1]].index("SalaryBasis") + 1)
        wb.save(self.path)
        loaded = TimecardBook.load(self.path)
        self.assertEqual(loaded.employees["E001"].salary_basis, "Annual")
        self.assertEqual(pay(loaded.employees["E001"], Company(pay_period="Weekly")), D("1000.00"))


if __name__ == "__main__":
    unittest.main()
