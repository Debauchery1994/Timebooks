"""Tests for timecard.storage workbook load/save."""
from __future__ import annotations

import os
import tempfile
import unittest
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path

import openpyxl

from timecard import calc
from timecard.models import Company, Employee, PayrollRecord, TimeEntry
from timecard.storage import TimecardBook, WorkbookError

D = Decimal


def sample_book() -> TimecardBook:
    company = Company(company_name="Acme Tile", address1="1 Main St", routing_number="011000015",
                      account_number="00123456", year=2026, next_check_number=2001,
                      pay_period="Biweekly", week_start="Sunday", default_lunch_minutes=45,
                      overtime_threshold_weekly=40.0, overtime_multiplier=2.0,
                      federal_withholding_pct=7.5, print_offset_x=0.125)
    book = TimecardBook.new(company)
    book.upsert_employee(Employee("E001", "Ann", "Lee", "5 Oak", "Town, ST 00001",
                                  pay_rate=D("21.50"), notes="lead"))
    book.upsert_employee(Employee("E002", "Bo", "Ray", pay_type="Salary", pay_rate=D("1500"),
                                  active=False))
    book.set_entry(TimeEntry(date(2026, 1, 5), "E001", time(8), time(12), time(12, 30), time(17),
                             15, "site A"))
    book.set_entry(TimeEntry(date(2026, 1, 6), "E001", time(22), time(6, 30)))
    book.set_entry(TimeEntry(date(2026, 1, 5), "E002", time(9), time(17), lunch_deduct_min=30))
    book.record_payroll([
        PayrollRecord(2001, date(2026, 1, 9), date(2026, 1, 4), date(2026, 1, 17), "E001", "Ann Lee",
                      regular_hours=D("16.25"), regular_pay=D("349.38"), gross_pay=D("349.38"),
                      social_security=D("21.66"), medicare=D("5.07"), net_pay=D("322.65"),
                      memo="first"),
        PayrollRecord(2002, date(2026, 1, 9), date(2026, 1, 4), date(2026, 1, 17), "E002", "Bo Ray",
                      gross_pay=D("1500.00"), net_pay=D("1385.25"), voided=True),
    ])
    return book


class StorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_round_trip_all_sheets(self) -> None:
        book = sample_book()
        path = book.save(self.dir / "acme.xlsx")
        self.assertFalse(book.dirty)
        self.assertEqual([p.name for p in self.dir.iterdir()], ["acme.xlsx"])  # temp file gone

        loaded = TimecardBook.load(path)
        self.assertEqual(loaded.company, book.company)
        self.assertEqual(loaded.employees, book.employees)
        self.assertEqual(loaded.entries, book.entries)
        self.assertEqual(loaded.payroll, book.payroll)
        self.assertEqual(loaded.path, path)
        self.assertFalse(loaded.dirty)
        self.assertEqual(loaded.company.next_check_number, 2003)

    def test_former_team_members_kept_for_bookkeeping(self) -> None:
        book = sample_book()  # E002 is inactive but has hours and a check this year
        book.upsert_employee(Employee("E003", "Cy", "Old", active=False))  # no history this year
        path = book.save(self.dir / "acme.xlsx")

        loaded = TimecardBook.load(path)
        self.assertIn("E003", loaded.employees)  # stored, never deleted
        self.assertEqual([e.employee_id for e in loaded.active_employees()], ["E001"])
        ids = [r[0].value for r in openpyxl.load_workbook(path)["Summary"].iter_rows(min_row=2, max_col=2)]
        names = [r[1].value for r in openpyxl.load_workbook(path)["Summary"].iter_rows(min_row=2, max_col=2)]
        self.assertIn("E002", ids)
        self.assertIn("Ray, Bo (former)", names)
        self.assertNotIn("E003", ids)

    def test_schema_and_formats(self) -> None:
        path = sample_book().save(self.dir / "acme.xlsx")
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames,
                         ["Company", "Employees", "TimeLog", "Payroll", "Summary", "Expenses", "Categories"])
        log = wb["TimeLog"]
        self.assertEqual([c.value for c in log[1]],
                         ["Date", "EmployeeID", "Employee", "In1", "Out1", "In2", "Out2",
                          "LunchDeductMin", "Hours", "Notes", "SalaryHours", "HoursType"])
        self.assertEqual(log.freeze_panes, "A2")
        self.assertTrue(log["A1"].font.bold)
        # Sorted by date then id; Hours is the live formula.
        self.assertEqual([(r[0].value.date(), r[1].value) for r in log.iter_rows(min_row=2)],
                         [(date(2026, 1, 5), "E001"), (date(2026, 1, 5), "E002"),
                          (date(2026, 1, 6), "E001")])
        self.assertTrue(log["I2"].value.startswith('=IF(AND(L2<>"",L2<>"Worked"),0,IF(K2<>"",'))
        self.assertIn('MAX(0,ROUND((IF(AND(D2<>"",E2<>""),MOD(E2-D2,1),0)', log["I2"].value)
        self.assertEqual(log["D2"].number_format, "h:mm AM/PM")
        self.assertEqual(log["C2"].value, "Lee, Ann")
        self.assertEqual(wb["Employees"]["G2"].number_format, "$#,##0.00")
        self.assertEqual(wb["Payroll"]["K2"].number_format, "$#,##0.00")
        self.assertEqual(wb["Payroll"]["S3"].value, "Yes")
        summary = wb["Summary"]
        self.assertEqual(summary["A2"].value, "E001")
        self.assertEqual(summary["B3"].value, "Ray, Bo (former)")  # former E002 kept: has history
        self.assertIn("SUMIFS(Payroll!K:K", summary["D2"].value)
        self.assertIn('"<>Yes"', summary["D2"].value)
        self.assertIn("DATE(2027,1,1)", summary["N6"].value)

    def test_load_tolerates_foreign_value_types(self) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Employees"
        ws.append(["EmployeeID", "FirstName", "LastName", "PayRate", "Active"])  # columns missing/moved
        ws.append(["E007", "Cy", "Dee", "$1,234.5", "no"])
        log = wb.create_sheet("TimeLog")
        log.append(["Date", "EmployeeID", "In1", "Out1", "In2", "Out2", "LunchDeductMin", "Hours"])
        log.append([46027, "E007", 0.375, "5:15 pm", datetime(1899, 12, 30, 18), time(19, 30), 30.0, 999])
        log.append(["2026-01-06", "E007", "8a", None, None, None, None, None])
        log.append([None, None, None, None, None, None, None, None])
        company = wb.create_sheet("Company")
        company.append(["Setting", "Value"])
        company.append(["company_name", "Foreign Co"])
        company.append(["year", 2026.0])
        company.append(["overtime_multiplier", "1.5"])
        company.append(["unknown_key", "ignored"])
        path = self.dir / "foreign.xlsx"
        wb.save(path)

        book = TimecardBook.load(path)
        self.assertEqual(book.company.company_name, "Foreign Co")
        self.assertEqual(book.company.year, 2026)
        self.assertEqual(book.company.medicare_pct, 0.0)  # default for missing key
        self.assertFalse(book.company.overtime_enabled)
        emp = book.employees["E007"]
        self.assertEqual(emp.pay_rate, D("1234.50"))
        self.assertFalse(emp.active)
        self.assertEqual(book.payroll, [])  # missing sheet treated as empty
        e = book.entries[("E007", date(2026, 1, 5))]
        self.assertEqual((e.in1, e.out1, e.in2, e.out2), (time(9), time(17, 15), time(18), time(19, 30)))
        self.assertEqual(e.lunch_deduct_min, 30)
        self.assertEqual(calc.day_hours(e), D("9.25"))
        self.assertEqual(book.entries[("E007", date(2026, 1, 6))].in1, time(8))

    def test_load_errors(self) -> None:
        with self.assertRaises(WorkbookError):
            TimecardBook.load(self.dir / "missing.xlsx")
        bad = self.dir / "bad.xlsx"
        bad.write_text("not a zip")
        with self.assertRaises(WorkbookError):
            TimecardBook.load(bad)
        other = self.dir / "other.xlsx"
        openpyxl.Workbook().save(other)
        with self.assertRaisesRegex(WorkbookError, "Employees or TimeLog"):
            TimecardBook.load(other)

    def test_save_without_path(self) -> None:
        with self.assertRaises(WorkbookError):
            TimecardBook.new(Company()).save()

    def test_save_permission_error_message(self) -> None:
        book = sample_book()
        path = book.save(self.dir / "acme.xlsx")
        real_replace = os.replace

        def locked(src, dst):  # simulate Windows sharing violation
            raise PermissionError(13, "Permission denied")

        os.replace = locked
        try:
            with self.assertRaisesRegex(WorkbookError, "open in another program"):
                book.save(path)
        finally:
            os.replace = real_replace
        self.assertEqual([p.name for p in self.dir.iterdir()], ["acme.xlsx"])  # temp cleaned up

    def test_upsert_preserves_pay_rate(self) -> None:
        book = sample_book()
        book.upsert_employee(Employee("E001", "Annie", "Lee", pay_rate=D("0")))
        self.assertEqual(book.employees["E001"].first_name, "Annie")
        self.assertEqual(book.employees["E001"].pay_rate, D("21.50"))
        book.upsert_employee(Employee("E001", "Annie", "Lee", pay_rate=D("99.99")))
        self.assertEqual(book.employees["E001"].pay_rate, D("21.50"))
        book.set_pay_rate("E001", D("23.255"))
        self.assertEqual(book.employees["E001"].pay_rate, D("23.26"))
        self.assertTrue(book.has_pay_rate("E001"))
        book.upsert_employee(Employee("E003", "Cy", "Dee"))
        self.assertFalse(book.has_pay_rate("E003"))
        self.assertEqual(book.next_employee_id(), "E004")

    def test_set_entry_empty_deletes(self) -> None:
        book = sample_book()
        d = date(2026, 1, 5)
        self.assertIn(("E001", d), book.entries)
        book.dirty = False
        book.set_entry(TimeEntry(d, "E001"))
        self.assertNotIn(("E001", d), book.entries)
        self.assertTrue(book.dirty)
        blank = book.get_entry("E001", d)
        self.assertTrue(blank.is_empty())
        self.assertEqual(blank.key, ("E001", d))

    def test_entries_and_employee_queries(self) -> None:
        book = sample_book()
        self.assertEqual([e.key for e in book.entries_for(None, date(2026, 1, 5), date(2026, 1, 6))],
                         [("E001", date(2026, 1, 5)), ("E002", date(2026, 1, 5)), ("E001", date(2026, 1, 6))])
        self.assertEqual(len(book.entries_for("E001", date(2026, 1, 6), date(2026, 1, 31))), 1)
        self.assertEqual([e.employee_id for e in book.active_employees()], ["E001"])
        # get_entry returns a copy: editing it does not change the book.
        copy = book.get_entry("E001", date(2026, 1, 5))
        copy.notes = "changed"
        self.assertEqual(book.entries[("E001", date(2026, 1, 5))].notes, "site A")

    def test_payroll_helpers(self) -> None:
        book = sample_book()
        book.void_check(2001)
        self.assertTrue(book.payroll[0].voided)
        with self.assertRaises(KeyError):
            book.void_check(42)
        recs = book.records_for_period(date(2026, 1, 4), date(2026, 1, 17))
        self.assertEqual([r.check_number for r in recs], [2001, 2002])
        self.assertEqual(book.records_for_period(date(2026, 1, 18), date(2026, 1, 31)), [])


if __name__ == "__main__":
    unittest.main()
