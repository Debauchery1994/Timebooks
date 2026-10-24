"""Tests for company branding: logo normalization, storage in the workbook, colors."""
from __future__ import annotations

import io
import random
import tempfile
import unittest
from pathlib import Path

import openpyxl
from PIL import Image

from timecard import branding
from timecard.models import Company
from timecard.storage import TimecardBook


def image_bytes(size=(1600, 900), fmt="PNG", noisy=False) -> bytes:
    img = Image.new("RGB", size, (30, 90, 160))
    if noisy:  # hard-to-compress pixels -> a big PNG that needs several cells
        rnd = random.Random(1)
        img.putdata([(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)) for _ in range(size[0] * size[1])])
    out = io.BytesIO()
    img.save(out, fmt)
    return out.getvalue()


class LogoTests(unittest.TestCase):
    def test_normalize_any_format_to_small_png(self) -> None:
        for fmt in ("PNG", "JPEG", "BMP", "GIF"):
            png = branding.normalize_logo(image_bytes(fmt=fmt))
            self.assertTrue(png.startswith(b"\x89PNG"), fmt)
            w, h = branding.logo_size(png)
            self.assertLessEqual(w, 800)
            self.assertLessEqual(h, 400)

    def test_transparent_border_trimmed(self) -> None:
        img = Image.new("RGBA", (300, 300), (0, 0, 0, 0))
        img.paste((200, 0, 0, 255), (100, 120, 200, 180))
        out = io.BytesIO()
        img.save(out, "PNG")
        self.assertEqual(branding.logo_size(branding.normalize_logo(out.getvalue())), (100, 60))

    def test_rejects_non_images(self) -> None:
        for bad in (b"", b"not an image at all", b"%PDF-1.4 ..."):
            with self.assertRaises(ValueError):
                branding.normalize_logo(bad)

    def test_fit_logo_keeps_shape(self) -> None:
        png = branding.normalize_logo(image_bytes((800, 400)))
        self.assertEqual(branding.fit_logo(png, 100, 100), (100, 50))
        self.assertEqual(branding.fit_logo(png, 400, 20), (40, 20))


class ColorTests(unittest.TestCase):
    def test_validation_and_default(self) -> None:
        self.assertTrue(branding.is_valid_color("#a1b2c3"))
        for bad in ("red", "#12345", "123456", "", "#GGGGGG"):
            self.assertFalse(branding.is_valid_color(bad))
            self.assertEqual(branding.brand_color(bad), branding.DEFAULT_BRAND_COLOR)
        self.assertEqual(branding.brand_color("#a1b2c3"), "#A1B2C3")

    def test_readable_text_colors(self) -> None:
        self.assertEqual(branding.text_on("#1F3A5F"), "#FFFFFF")
        self.assertEqual(branding.text_on("#FFE066"), "#000000")
        self.assertEqual(branding.ink_color("#1F3A5F"), "#1F3A5F")  # dark enough for print
        self.assertEqual(branding.ink_color("#FFE066"), "#222222")  # too light for text on paper
        self.assertGreater(branding.contrast_ratio("#000000", "#FFFFFF"), 20)


class WorkbookLogoTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name)

    def test_logo_round_trips_in_hidden_sheet(self) -> None:
        book = TimecardBook.new(Company(company_name="Acme", brand_color="#0B6E4F", logo_on_reports=False))
        logo_file = self.dir / "logo.jpg"
        logo_file.write_bytes(image_bytes(fmt="JPEG"))
        book.set_logo(logo_file)
        path = book.save(self.dir / "acme.xlsx")

        loaded = TimecardBook.load(path)
        self.assertEqual(loaded.logo_png, book.logo_png)
        self.assertEqual(loaded.company.brand_color, "#0B6E4F")
        self.assertFalse(loaded.company.logo_on_reports)
        self.assertTrue(loaded.company.logo_on_checks)
        ws = openpyxl.load_workbook(path)["Branding"]
        self.assertEqual(ws.sheet_state, "hidden")

    def test_large_logo_split_across_cells(self) -> None:
        book = TimecardBook.new(Company())
        book.set_logo(image_bytes((800, 400), noisy=True))
        self.assertGreater(len(book.logo_png) * 4 / 3, 32_767)  # too big for one cell
        path = book.save(self.dir / "big.xlsx")
        rows = list(openpyxl.load_workbook(path)["Branding"].iter_rows(min_row=2, values_only=True))
        self.assertGreater(len(rows), 1)
        self.assertTrue(all(len(v) <= 32_767 for _, v in rows))
        self.assertEqual(TimecardBook.load(path).logo_png, book.logo_png)

    def test_no_logo_means_no_branding_sheet_and_clear_works(self) -> None:
        book = TimecardBook.new(Company())
        self.assertNotIn("Branding", openpyxl.load_workbook(book.save(self.dir / "plain.xlsx")).sheetnames)
        book.set_logo(image_bytes())
        book.clear_logo()
        path = book.save(self.dir / "cleared.xlsx")
        self.assertEqual(TimecardBook.load(path).logo_png, b"")

    def test_damaged_logo_does_not_block_opening(self) -> None:
        book = TimecardBook.new(Company())
        book.set_logo(image_bytes())
        path = book.save(self.dir / "acme.xlsx")
        wb = openpyxl.load_workbook(path)
        wb["Branding"]["B2"] = "!!! not base64 !!!"
        wb.save(path)
        self.assertEqual(TimecardBook.load(path).logo_png, b"")


if __name__ == "__main__":
    unittest.main()
