"""Workbook schema version: older workbooks upgrade on save; a newer app's workbook is never saved over."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import openpyxl

from timecard import __version__
from timecard.models import Company, Employee
from timecard.storage import SCHEMA_VERSION, TimecardBook, WorkbookError


def _company_rows(path: Path) -> dict:
    ws = openpyxl.load_workbook(path)["Company"]
    return {r[0]: r[1] for r in ws.iter_rows(min_row=2, max_col=2, values_only=True)}


def _set_meta(path: Path, version: int | None) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb["Company"]
    for row in list(ws.iter_rows(min_row=2)):
        if row[0].value in ("schema_version", "app_version"):
            if version is None:
                ws.delete_rows(row[0].row)
            elif row[0].value == "schema_version":
                row[1].value = version
            else:
                row[1].value = "9.9.9"
    wb.save(path)


class SchemaVersionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "b.xlsx"
        b = TimecardBook.new(Company(year=2026, company_name="Acme"))
        b.upsert_employee(Employee("E1", "Ann", "Lee"))
        b.save(self.path)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_save_writes_the_marker(self) -> None:
        rows = _company_rows(self.path)
        self.assertEqual((rows["schema_version"], rows["app_version"]), (SCHEMA_VERSION, __version__))
        book = TimecardBook.load(self.path)
        self.assertEqual(book.schema_version, SCHEMA_VERSION)
        self.assertFalse(book.from_newer_app)
        self.assertEqual(book.company.company_name, "Acme")

    def test_older_workbook_without_marker_loads_and_upgrades(self) -> None:
        _set_meta(self.path, None)
        book = TimecardBook.load(self.path)
        self.assertEqual(book.schema_version, 1)
        self.assertFalse(book.from_newer_app)
        book.save()
        self.assertEqual(_company_rows(self.path)["schema_version"], SCHEMA_VERSION)

    def test_newer_workbook_is_never_saved_over(self) -> None:
        _set_meta(self.path, SCHEMA_VERSION + 1)
        before = self.path.read_bytes()
        book = TimecardBook.load(self.path)
        self.assertTrue(book.from_newer_app)
        self.assertIn("9.9.9", book.newer_message())
        book.dirty = True
        with self.assertRaises(WorkbookError):
            book.save()
        with self.assertRaises(WorkbookError):
            book.save(str(self.path))
        self.assertEqual(self.path.read_bytes(), before)
        copy = book.save(Path(self.tmp.name) / "copy.xlsx")   # Save As a separate copy is allowed
        self.assertEqual(_company_rows(copy)["schema_version"], SCHEMA_VERSION)
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
