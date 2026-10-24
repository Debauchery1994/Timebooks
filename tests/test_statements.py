"""Tests for annual statements export."""
from __future__ import annotations

import tempfile
import unittest
from datetime import date, time
from decimal import Decimal
from pathlib import Path

import openpyxl

from timecard import statements
from timecard.models import Company, Employee, PayrollRecord, TimeEntry
from timecard.storage import TimecardBook

D = Decimal


def check(n: int, d: date, emp: str, gross: str, other: str = "0", voided: bool = False, hours: str = "40") -> PayrollRecord:
    g, o = D(gross), D(other)
    return PayrollRecord(n, d, d, d, emp, emp, regular_hours=D(hours), regular_pay=g, gross_pay=g,
                         other_deductions=o, net_pay=g - o, voided=voided)


def sample_book() -> TimecardBook:
    book = TimecardBook.new(Company(company_name="Acme/Tile", year=2026))
    book.upsert_employee(Employee("E001", "Ann", "Lee", "5 Oak", "Town, ST 00001"))
    book.set_pay_rate("E001", D("23.47"))
    book.upsert_employee(Employee("E002", "Bo", "Ray", active=False))  # former, paid this year
    book.set_pay_rate("E002", D("23.47"))
    book.upsert_employee(Employee("E003", "Cy", "Idle"))  # nothing this year
    book.set_entry(TimeEntry(date(2026, 1, 5), "E001", time(8), time(17), lunch_deduct_min=30))
    book.set_entry(TimeEntry(date(2025, 12, 31), "E001", time(8), time(12)))  # other year: ignored
    book.payroll = [
        check(1001, date(2026, 1, 9), "E001", "1500.00", other="-25.00"),  # +25 reimbursement
        check(1002, date(2026, 3, 6), "E001", "900.00", other="50.00"),
        check(1003, date(2026, 3, 6), "E002", "700.00"),
        check(1004, date(2026, 4, 3), "E002", "999.00", voided=True),
        check(1005, date(2025, 12, 31), "E001", "5000.00"),  # prior year: ignored
    ]
    return book


class StatementTests(unittest.TestCase):
    def test_totals_exclude_voided_and_other_years(self) -> None:
        by_id = {s.employee.employee_id: s for s in statements.build_statements(sample_book())}
        self.assertEqual(set(by_id), {"E001", "E002"})  # E003 had no activity
        ann = by_id["E001"]
        self.assertEqual(ann.gross, D("2400.00"))
        self.assertEqual(ann.adjustments, D("-25.00"))  # +25 added, -50 deducted
        self.assertEqual(ann.paid, D("2375.00"))
        self.assertEqual(ann.hours, D("8.50"))
        self.assertEqual(ann.monthly_paid[0], D("1525.00"))
        self.assertEqual(ann.monthly_paid[2], D("850.00"))
        self.assertTrue(ann.at_or_over(D("2000")))
        bo = by_id["E002"]
        self.assertEqual(bo.paid, D("700.00"))
        self.assertFalse(bo.at_or_over(D("2000")))

    def test_filter_by_employee(self) -> None:
        only = statements.build_statements(sample_book(), ["E002"])
        self.assertEqual([s.employee.employee_id for s in only], ["E002"])

    def test_export_pdf_and_xlsx_without_pay_rate(self) -> None:
        book = sample_book()
        with tempfile.TemporaryDirectory() as tmp:
            pdf, xlsx = statements.export_annual_statements(book, tmp)
            self.assertEqual(pdf.name, "AcmeTile Annual Statements 2026.pdf")
            self.assertGreater(pdf.stat().st_size, 2000)
            wb = openpyxl.load_workbook(xlsx)
            self.assertEqual(wb.sheetnames, ["Summary", "Payments", "Monthly"])
            summary = list(wb["Summary"].iter_rows(values_only=True))
            self.assertEqual(summary[1][:2], ("E001", "Ann Lee"))
            self.assertEqual(summary[1][8], 2375.0)
            self.assertEqual(summary[1][10], "Yes")
            self.assertEqual(summary[2][4], "Former")
            self.assertEqual(len(list(wb["Payments"].iter_rows(min_row=2))), 3)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    self.assertNotIn(23.47, row)
                    self.assertNotIn("23.47", [str(v) for v in row])

            one_pdf, _ = statements.export_annual_statements(book, tmp, ["E002"])
            self.assertEqual(one_pdf.name, "AcmeTile Annual Statements 2026 - Bo Ray.pdf")

    def test_pdf_text_has_no_pay_rate(self) -> None:
        from reportlab import rl_config
        old = rl_config.pageCompression
        rl_config.pageCompression = 0
        try:
            with tempfile.TemporaryDirectory() as tmp:
                pdf, _ = statements.export_annual_statements(sample_book(), tmp)
                data = pdf.read_bytes()
        finally:
            rl_config.pageCompression = old
        self.assertIn(b"Annual Payment Summary", data)
        self.assertNotIn(b"23.47", data)

    def test_no_activity_raises(self) -> None:
        book = TimecardBook.new(Company(year=2026))
        with tempfile.TemporaryDirectory() as tmp, self.assertRaises(ValueError):
            statements.export_annual_statements(book, tmp)


def _logo_png() -> bytes:
    import io
    from PIL import Image, ImageDraw
    img = Image.new("RGBA", (300, 100), (0, 0, 0, 0))
    ImageDraw.Draw(img).ellipse((5, 5, 95, 95), fill=(0, 128, 128, 255))
    out = io.BytesIO()
    img.save(out, "PNG")
    return out.getvalue()


class StatementBrandingTests(unittest.TestCase):
    def _pdf_bytes(self, book) -> bytes:
        from reportlab import rl_config
        old = rl_config.pageCompression
        rl_config.pageCompression = 0
        try:
            with tempfile.TemporaryDirectory() as tmp:
                pdf, _ = statements.export_annual_statements(book, tmp)
                return pdf.read_bytes()
        finally:
            rl_config.pageCompression = old

    def test_logo_and_brand_color(self) -> None:
        from reportlab.lib.colors import HexColor
        from reportlab.lib.rl_accel import fp_str
        book = sample_book()
        book.logo_png = _logo_png()
        book.company.brand_color = "#00807F"
        data = self._pdf_bytes(book)
        self.assertIn(b"/Subtype /Image", data)
        c = HexColor("#00807F")
        self.assertIn(fp_str(c.red, c.green, c.blue).encode(), data)
        self.assertNotIn(b"23.47", data)

    def test_logo_off_or_missing(self) -> None:
        book = sample_book()
        book.logo_png = _logo_png()
        book.company.logo_on_reports = False
        self.assertNotIn(b"/Subtype /Image", self._pdf_bytes(book))
        book = sample_book()
        book.company.brand_color = "not a color"
        data = self._pdf_bytes(book)
        self.assertNotIn(b"/Subtype /Image", data)
        self.assertIn(b"Annual Payment Summary", data)


if __name__ == "__main__":
    unittest.main()
