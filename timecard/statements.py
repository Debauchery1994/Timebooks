"""Annual payment statements for a tax professional (PDF + Excel).

A payment belongs to the year of its check date (cash basis, as for Form 1099-NEC). Voided checks
are left out. Pay rates are never included. Former team members are included when they were paid
or logged hours during the year.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from . import calc
from .models import Employee, PayrollRecord

if TYPE_CHECKING:
    from .storage import TimecardBook

ZERO = Decimal("0")
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONEY_FMT = "$#,##0.00"
SHADE = colors.Color(0.92, 0.92, 0.92)


@dataclass
class AnnualStatement:
    employee: Employee
    hours: Decimal = ZERO
    gross: Decimal = ZERO          # compensation for work (hours x rate, or salary)
    adjustments: Decimal = ZERO    # effect on checks: + added (e.g. reimbursement), - deducted
    paid: Decimal = ZERO           # total of check amounts
    monthly_paid: list[Decimal] = field(default_factory=lambda: [ZERO] * 12)
    checks: list[PayrollRecord] = field(default_factory=list)

    def at_or_over(self, threshold: Decimal) -> bool:
        return self.gross >= threshold


def build_statements(book: "TimecardBook", employee_ids: list[str] | None = None) -> list[AnnualStatement]:
    """Statements for everyone paid or with hours in book.company.year (optionally only `employee_ids`)."""
    year = book.company.year
    start, end = date(year, 1, 1), date(year, 12, 31)
    by_id: dict[str, AnnualStatement] = {}

    def stmt(emp_id: str) -> AnnualStatement | None:
        if employee_ids is not None and emp_id not in employee_ids:
            return None
        if emp_id not in by_id:
            emp = book.employees.get(emp_id) or Employee(emp_id, "", emp_id, active=False)
            by_id[emp_id] = AnnualStatement(emp)
        return by_id[emp_id]

    for e in book.entries.values():
        if start <= e.work_date <= end and (s := stmt(e.employee_id)):
            s.hours += calc.day_hours(e)
    for r in sorted(book.payroll, key=lambda r: (r.check_date, r.check_number)):
        if r.voided or not start <= r.check_date <= end or not (s := stmt(r.employee_id)):
            continue
        s.checks.append(r)
        s.gross += r.gross_pay
        s.adjustments -= r.other_deductions
        s.paid += r.net_pay
        s.monthly_paid[r.check_date.month - 1] += r.net_pay
    return sorted(by_id.values(), key=lambda s: (s.employee.last_name.lower(), s.employee.first_name.lower()))


def _money(x: Decimal) -> str:
    return f"${calc.money(x):,.2f}"


def _signed(x: Decimal) -> str:
    x = calc.money(x)
    return "—" if x == 0 else (f"+{x:,.2f}" if x > 0 else f"{x:,.2f}")


def _name(emp: Employee) -> str:
    return emp.full_name + ("" if emp.active else " (former)")


def _safe(s: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "", s).strip() or "Company"


def default_basename(book: "TimecardBook") -> str:
    return f"{_safe(book.company.company_name)} Annual Statements {book.company.year}"


# ---------------------------------------------------------------- PDF

def _table(rows: list[list], widths: list[float], money_cols: tuple[int, ...] = (), total_row: bool = False) -> Table:
    t = Table(rows, colWidths=widths, repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8.5),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
        ("BACKGROUND", (0, 0), (-1, 0), SHADE),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for col in money_cols:
        style.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
    if total_row:
        style += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8.5),
                  ("LINEABOVE", (0, -1), (-1, -1), 0.6, colors.black)]
    t.setStyle(TableStyle(style))
    return t


def render_statements_pdf(book: "TimecardBook", statements: list[AnnualStatement], out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    co = book.company
    year = co.year
    threshold = Decimal(str(co.form_1099_threshold))
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=16, alignment=0, spaceAfter=2)
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, leading=11)
    note = ParagraphStyle("note", parent=small, textColor=colors.Color(0.35, 0.35, 0.35))
    company_lines = "<br/>".join(x for x in (co.company_name, co.address1, co.address2, co.phone) if x)
    printed = datetime.now().strftime("%m/%d/%Y")

    def header(title: str) -> list:
        return [Paragraph(title, h1), Paragraph(company_lines, small),
                Paragraph(f"Calendar year {year} &middot; prepared {printed}", note), Spacer(1, 10)]

    story: list = header(f"Annual Payment Summary — {year}")
    rows = [["ID", "Name", "Hours", "Compensation", "Adjustments", "Total Paid", "Checks", "1099 review"]]
    for s in statements:
        rows.append([s.employee.employee_id, _name(s.employee), f"{s.hours:,.2f}", _money(s.gross),
                     _signed(s.adjustments), _money(s.paid), str(len(s.checks)),
                     "At/over threshold" if s.at_or_over(threshold) else ""])
    rows.append(["", "Total", f"{sum((s.hours for s in statements), ZERO):,.2f}",
                 _money(sum((s.gross for s in statements), ZERO)),
                 _signed(sum((s.adjustments for s in statements), ZERO)),
                 _money(sum((s.paid for s in statements), ZERO)),
                 str(sum(len(s.checks) for s in statements)), ""])
    story.append(_table(rows, [0.5 * inch, 1.9 * inch, 0.6 * inch, 0.95 * inch, 0.85 * inch, 0.9 * inch,
                               0.5 * inch, 1.05 * inch], money_cols=(2, 3, 4, 5, 6), total_row=True))
    story += [Spacer(1, 10), Paragraph(
        f"Compensation is pay for work (hours &times; rate, or salary). Adjustments are amounts added to (+) or "
        f"deducted from (&minus;) checks, such as reimbursements. Payments are counted in the year of the check "
        f"date; voided checks are excluded. \"At/over threshold\" marks compensation of "
        f"{_money(threshold)} or more (the Form 1099-NEC reporting threshold set in Company Settings). "
        f"Taxpayer ID numbers are not kept in this program — collect a Form W-9 from each payee. "
        f"Please confirm reporting requirements with your tax professional.", note)]

    for s in statements:
        emp = s.employee
        story += [PageBreak(), *header(f"Annual Statement — {year}")]
        addr = "<br/>".join(x for x in (f"<b>{_name(emp)}</b>", emp.address1, emp.address2) if x)
        info = Table([[Paragraph(addr, small), Paragraph(
            f"Payee ID: {emp.employee_id}<br/>Pay type: {'Salaried' if emp.pay_type.lower() == 'salary' else 'Hourly'}"
            f"<br/>Status: {'Active' if emp.active else 'No longer with company'}", small)]],
            colWidths=[3.6 * inch, 3.6 * inch])
        info.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        story += [info, Spacer(1, 8)]
        totals = [["Hours worked", "Compensation", "Adjustments", "Total paid", "Checks"],
                  [f"{s.hours:,.2f}", _money(s.gross), _signed(s.adjustments), _money(s.paid), str(len(s.checks))]]
        story += [_table(totals, [1.44 * inch] * 5, money_cols=(0, 1, 2, 3, 4)), Spacer(1, 10)]
        cells = [_money(m) if m else "—" for m in s.monthly_paid]
        mt = _table([MONTHS[:6], cells[:6], MONTHS[6:], cells[6:]], [1.2 * inch] * 6, money_cols=tuple(range(6)))
        mt.setStyle(TableStyle([("FONT", (0, 2), (-1, 2), "Helvetica-Bold", 8.5), ("BACKGROUND", (0, 2), (-1, 2), SHADE),
                                ("LINEBELOW", (0, 2), (-1, 2), 0.6, colors.black)]))
        story += [Paragraph("<b>Paid by month</b>", small), Spacer(1, 3), mt, Spacer(1, 10)]
        crow = [["Check #", "Date", "Pay period", "Hours", "Compensation", "Adjustment", "Check amount"]]
        for r in s.checks:
            crow.append([str(r.check_number), r.check_date.strftime("%m/%d/%Y"),
                         f"{r.period_start:%m/%d}–{r.period_end:%m/%d/%Y}",
                         f"{r.regular_hours + r.overtime_hours:,.2f}", _money(r.gross_pay),
                         _signed(-r.other_deductions), _money(r.net_pay)])
        crow.append(["", "", "Total", f"{sum((r.regular_hours + r.overtime_hours for r in s.checks), ZERO):,.2f}",
                     _money(s.gross), _signed(s.adjustments), _money(s.paid)])
        story += [Paragraph("<b>Checks issued</b>", small), Spacer(1, 3),
                  _table(crow, [0.7 * inch, 0.85 * inch, 1.6 * inch, 0.7 * inch, 1.1 * inch, 1.0 * inch, 1.1 * inch],
                         money_cols=(3, 4, 5, 6), total_row=True)]
        if not s.checks:
            story.append(Paragraph("No checks were issued to this payee during the year.", note))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.grey)
        canvas.drawString(0.6 * inch, 0.45 * inch, f"{co.company_name} — {year} annual statements — confidential")
        canvas.drawRightString(letter[0] - 0.6 * inch, 0.45 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(out), pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.7 * inch,
                            title=f"{co.company_name} Annual Statements {year}", author="TimeCard Pro")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out


# ---------------------------------------------------------------- Excel

def _sheet(ws, headers: list[str], widths: list[int]) -> None:
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        ws.cell(row=1, column=i).font = Font(bold=True)
    ws.freeze_panes = "A2"


def export_statements_xlsx(book: "TimecardBook", statements: list[AnnualStatement], out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    threshold = Decimal(str(book.company.form_1099_threshold))
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    _sheet(ws, ["PayeeID", "Name", "Address1", "Address2", "Status", "Hours", "Compensation", "Adjustments",
                "TotalPaid", "Checks", "AtOrOver1099Threshold"], [10, 26, 28, 26, 10, 10, 14, 13, 14, 8, 12])
    for s in statements:
        e = s.employee
        ws.append([e.employee_id, e.full_name, e.address1, e.address2, "Active" if e.active else "Former",
                   float(s.hours), float(s.gross), float(s.adjustments), float(s.paid), len(s.checks),
                   "Yes" if s.at_or_over(threshold) else "No"])
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=9):
        for c in row:
            c.number_format = MONEY_FMT

    wp = wb.create_sheet("Payments")
    _sheet(wp, ["PayeeID", "Name", "CheckNumber", "CheckDate", "PeriodStart", "PeriodEnd", "Hours",
                "Compensation", "Adjustment", "CheckAmount", "Memo"], [10, 26, 11, 12, 12, 12, 9, 14, 12, 14, 40])
    for s in statements:
        for r in s.checks:
            wp.append([r.employee_id, s.employee.full_name, r.check_number, r.check_date, r.period_start,
                       r.period_end, float(r.regular_hours + r.overtime_hours), float(r.gross_pay),
                       float(-r.other_deductions), float(r.net_pay), r.memo])
    for row in wp.iter_rows(min_row=2):
        for c in row[3:6]:
            c.number_format = "yyyy-mm-dd"
        for c in row[7:10]:
            c.number_format = MONEY_FMT

    wm = wb.create_sheet("Monthly")
    _sheet(wm, ["PayeeID", "Name", *MONTHS, "Total"], [10, 26] + [11] * 13)
    for s in statements:
        wm.append([s.employee.employee_id, s.employee.full_name, *[float(m) for m in s.monthly_paid], float(s.paid)])
    for row in wm.iter_rows(min_row=2, min_col=3):
        for c in row:
            c.number_format = MONEY_FMT

    wb.properties.title = f"{book.company.company_name} Annual Statements {book.company.year}"
    wb.save(out)
    return out


def export_annual_statements(book: "TimecardBook", out_dir: str | Path,
                             employee_ids: list[str] | None = None) -> tuple[Path, Path]:
    """Write '<Company> Annual Statements <year>.pdf' and '.xlsx' to out_dir; returns both paths."""
    statements = build_statements(book, employee_ids)
    if not statements:
        raise ValueError(f"No hours or payments were recorded in {book.company.year}.")
    base = Path(out_dir) / default_basename(book)
    if employee_ids is not None and len(statements) == 1:
        base = base.with_name(f"{base.name} - {_safe(statements[0].employee.full_name)}")
    return (render_statements_pdf(book, statements, base.with_suffix(".pdf")),
            export_statements_xlsx(book, statements, base.with_suffix(".xlsx")))
